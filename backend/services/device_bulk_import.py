"""장비를 CSV/Excel 파일로 한 번에 여러 개 등록하는 기능 — 개별 등록 화면과 같은
검증(`routers/devices.py`의 `_validation_error`)을 행 단위로 재사용하고, 한 행이
잘못돼도 나머지 행은 계속 등록되도록 행 단위로 성공/실패를 모아 반환한다.

`file_extract.py`(App 1 등에서 쓰는 "문서를 AI 분석용 텍스트로" 추출기)와는 목적이
다르다 — 이건 표 형식 데이터를 "장비 등록 필드"로 파싱하는 것이라 별도 모듈로 둔다.
업로드된 파일 자체는 메모리에서만 처리하고 어디에도 저장하지 않는다(비밀번호가
포함될 수 있는 파일이라 이 프로젝트의 "원본 미저장" 원칙을 특히 더 지켜야 함).
"""

import csv
import io

from openpyxl import load_workbook

COLUMNS = ["name", "device_type", "vendor", "analyzer", "host", "port", "username", "password", "context", "interval_minutes", "enabled"]


def _normalize_header(raw: str) -> str:
    """헤더 셀에 공백/대소문자/하이픈이 섞여 있어도(" Device Type", "device-type") 표준
    컬럼명("device_type")으로 인식하도록 정규화한다."""
    return (raw or "").strip().lower().replace(" ", "_").replace("-", "_")


class ParseError(ValueError):
    pass


def parse_rows(filename: str, raw: bytes) -> list[dict]:
    """파일에서 행을 읽어 {컬럼명: 원본 값} 딕셔너리 리스트로 반환한다(타입 변환은
    `normalize_row`에서 별도로 함) — 값이 하나도 없는 완전히 빈 행은 건너뛴다."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        if reader.fieldnames is None:
            raise ParseError("CSV 헤더 행을 찾을 수 없습니다.")
        rows = []
        for row in reader:
            normalized = {_normalize_header(k): (v or "").strip() for k, v in row.items() if k}
            if any(v for v in normalized.values()):
                rows.append(normalized)
        return rows

    if name.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        except Exception as e:
            raise ParseError(f"Excel 파일을 읽을 수 없습니다: {e}")
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            return []
        header = [_normalize_header(str(h)) if h is not None else "" for h in all_rows[0]]
        rows = []
        for r in all_rows[1:]:
            if all(c is None or str(c).strip() == "" for c in r):
                continue
            row = {header[i]: r[i] for i in range(min(len(header), len(r))) if header[i]}
            rows.append(row)
        return rows

    raise ParseError("지원하지 않는 파일 형식입니다 — .csv 또는 .xlsx 파일을 업로드하세요.")


def _to_bool(value, default: bool = True) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("true", "1", "yes", "y", "on", "참", "예")


def _to_int(value, default: int | None):
    if value is None or value == "":
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        raise ParseError(f"숫자로 변환할 수 없는 값입니다: {value!r}")


def normalize_row(row: dict) -> dict:
    """엑셀/CSV에서 읽은 문자열/원시값을 `DevicePayload`가 기대하는 타입으로 변환한다.
    변환 실패는 ParseError로 올려 라우터가 그 행만 오류 목록에 담고 나머지는 계속
    처리하게 한다."""

    def s(key: str) -> str:
        v = row.get(key)
        return "" if v is None else str(v).strip()

    payload = {
        "name": s("name"),
        "device_type": s("device_type"),
        "vendor": s("vendor") or None,
        "analyzer": s("analyzer"),
        "host": s("host"),
        "port": _to_int(row.get("port"), None),
        "username": s("username"),
        "password": s("password") or None,
        "context": s("context"),
        "interval_minutes": _to_int(row.get("interval_minutes"), 60),
        "enabled": _to_bool(row.get("enabled"), True),
    }
    return payload
